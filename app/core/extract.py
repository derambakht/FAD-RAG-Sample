"""استخراج متن از PDF، DOCX، TXT و MD (بخش‌های ۳ و ۴ راهنما).

خروجی، فهرستی از خطوط همراه با شماره صفحه و اندازه فونت است تا مرحله بخش‌بندی
بتواند عنوان‌ها را از روی اندازه فونت تشخیص دهد و هر chunk شماره صفحه داشته باشد.
"""
from __future__ import annotations

import io
import logging
import re
import shutil
from collections import Counter
from dataclasses import asdict, dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Callable

from app.core.spreadsheet import TABLE_EXT
from app.core.normalize import char_stats, clean_fa, looks_broken, pre_clean, unreverse_ltr_runs
from app.core.text_utils import fa_token_count, vocab_score

log = logging.getLogger(__name__)

SUPPORTED_EXT = {".pdf", ".docx", ".txt", ".md"} | TABLE_EXT  # جدول‌ها در app.core.spreadsheet خوانده می‌شوند
Progress = Callable[[float, str], None]


@dataclass
class Line:
    text: str
    page: int
    size: float = 0.0
    bold: bool = False
    heading: bool = False  # عنوان صریح (سبک Heading در DOCX یا # در Markdown)


@dataclass
class PageResult:
    page: int
    lines: list[Line]
    method: str
    vocab: float = 0.0
    reversed: bool = False
    stats: dict = field(default_factory=dict)

    def summary(self) -> dict:
        return {"page": self.page, "method": self.method, "lines": len(self.lines),
                "vocab": round(self.vocab, 3), "reversed": self.reversed, **self.stats}


@dataclass
class ExtractOptions:
    method: str = "auto"  # auto | pymupdf | pdfplumber | ocr
    clean: bool = True  # اعمال pre_clean روی خطوط (برای آزمایش «بدون clean_fa» خاموش شود)
    margin: float = 0.06  # نوار سرصفحه و پاورقی بر اساس ارتفاع صفحه
    fix_direction: bool = True  # تشخیص و اصلاح خطوط با حروف معکوس
    reverse_word_order: bool = False  # کلمات درست اما ترتیبشان معکوس (تصمیم دستی برای کل سند)
    remove_repeated: bool = True  # حذف سرصفحه/پاورقی تکرارشونده و شماره صفحه
    extract_tables: bool = False  # جدول‌ها به markdown و chunk جداگانه
    ocr_fallback: bool = True

    @classmethod
    def from_dict(cls, d: dict | None) -> "ExtractOptions":
        d = d or {}
        return cls(**{k: v for k, v in d.items() if k in cls.__dataclass_fields__})


@dataclass
class ExtractResult:
    pages: list[PageResult]
    tables: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def lines(self) -> list[Line]:
        return [ln for p in self.pages for ln in p.lines]

    def summary(self) -> dict:
        methods = Counter(p.method for p in self.pages)
        return {
            "pages": len(self.pages),
            "methods": dict(methods),
            "reversed_pages": sum(p.reversed for p in self.pages),
            "avg_vocab": round(sum(p.vocab for p in self.pages) / max(len(self.pages), 1), 3),
            "tables": len(self.tables),
            "warnings": self.warnings,
        }


# --- OCR ----------------------------------------------------------------------------
@lru_cache(maxsize=1)
def ocr_status() -> dict:
    from app.config import get_settings
    s = get_settings()
    try:
        import pytesseract
    except ImportError:
        return {"available": False, "reason": "بسته pytesseract نصب نیست"}
    if s.tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = s.tesseract_cmd
    elif not shutil.which("tesseract"):
        return {"available": False, "reason": "برنامه Tesseract پیدا نشد"}
    try:
        langs = pytesseract.get_languages()
    except Exception as e:  # pragma: no cover
        return {"available": False, "reason": str(e)}
    if s.ocr_lang not in langs:
        return {"available": False, "reason": f"زبان {s.ocr_lang} برای Tesseract نصب نیست"}
    return {"available": True, "langs": langs}


def _ocr_lines(fpage, pno: int) -> list[Line]:
    import pytesseract
    from PIL import Image

    from app.config import get_settings
    pix = fpage.get_pixmap(dpi=300)
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    text = pytesseract.image_to_string(img, lang=get_settings().ocr_lang)
    return [Line(t, pno) for t in text.splitlines() if t.strip()]


# --- PyMuPDF ------------------------------------------------------------------------
def _is_rtl(text: str) -> bool:
    rtl = sum("؀" <= c <= "ۿ" or "ﭐ" <= c <= "﻿" for c in text)
    ltr = sum(c.isascii() and c.isalpha() for c in text)
    return rtl >= ltr


def _pymupdf_lines(fpage, pno: int, margin: float) -> list[Line]:
    h = fpage.rect.height
    out: list[Line] = []
    for b in fpage.get_text("dict", sort=True)["blocks"]:
        if b["type"] != 0:
            continue
        for line in b["lines"]:
            y0, y1 = line["bbox"][1], line["bbox"][3]
            if margin and (y1 < margin * h or y0 > (1 - margin) * h):  # سرصفحه و پاورقی
                continue
            spans = [s for s in line["spans"] if s["text"]]
            if not spans:
                continue
            raw = "".join(s["text"] for s in spans)
            if _is_rtl(raw) and len(spans) > 1:  # ترتیب خواندن راست به چپ
                spans.sort(key=lambda s: -s["bbox"][0])
            parts = [spans[0]["text"]]
            for prev, cur in zip(spans, spans[1:]):
                gap = abs(prev["bbox"][0] - cur["bbox"][2]) if _is_rtl(raw) else abs(cur["bbox"][0] - prev["bbox"][2])
                if gap > 0.2 * cur["size"] and not parts[-1].endswith(" ") and not cur["text"].startswith(" "):
                    parts.append(" ")
                parts.append(cur["text"])
            text = "".join(parts)
            if not text.strip():
                continue
            out.append(Line(text=text, page=pno, size=round(max(s["size"] for s in spans), 1),
                            bold=all(s["flags"] & 16 for s in spans)))
    return out


# --- pdfplumber ---------------------------------------------------------------------
TOLERANCES = (0.8, 1.2, 1.5, 2.0, 3.0)


def _plumber_lines(ppage, pno: int, margin: float, tol: float) -> list[Line]:
    h = ppage.height
    out = []
    for ln in ppage.extract_text_lines(x_tolerance=tol, return_chars=True, strip=True):
        if margin and (ln["bottom"] < margin * h or ln["top"] > (1 - margin) * h):
            continue
        if not ln["text"].strip():
            continue
        size = max((c.get("size", 0) for c in ln.get("chars", [])), default=0)
        out.append(Line(text=ln["text"], page=pno, size=round(size, 1)))
    return out


def _score_lines(lines: list[Line]) -> float:
    """امتیاز واژگان با در نظر گرفتن احتمال معکوس بودن خطوط."""
    text = pre_clean("\n".join(ln.text for ln in lines))
    return max(vocab_score(text), vocab_score(text[::-1]))


def _best_plumber(ppage, pno: int, margin: float) -> tuple[list[Line], float]:
    best: tuple[list[Line], float] = ([], -1.0)
    for tol in TOLERANCES:
        lines = _plumber_lines(ppage, pno, margin, tol)
        sc = _score_lines(lines)
        if sc > best[1]:
            best = (lines, sc)
    return best


# --- جهت خط -------------------------------------------------------------------------
def reverse_line(text: str) -> str:
    return unreverse_ltr_runs(text[::-1])


def fix_line_direction(line: str) -> tuple[str, bool]:
    """انتخاب بین حالت عادی و حروف معکوس با امتیاز واژگان (برای تک‌خط)."""
    if fa_token_count(line) < 2:
        return line, False
    s0, rev = vocab_score(line), reverse_line(line)
    if vocab_score(rev) > s0 + 0.15:
        return rev, True
    return line, False


def _fix_page_direction(lines: list[Line]) -> bool:
    """تصمیم در سطح صفحه: اگر بیشتر خطوط با معکوس شدن بهتر شوند، همه خطوط فارسی معکوس می‌شوند.

    unreverse_ltr_runs فقط روی خطوطی اجرا می‌شود که واقعاً معکوس شده‌اند (هشدار بخش ۵ راهنما).
    """
    votes = better = 0
    for ln in lines:
        if fa_token_count(ln.text) < 3:
            continue
        votes += 1
        if vocab_score(reverse_line(ln.text)) > vocab_score(ln.text) + 0.1:
            better += 1
    if votes == 0 or better / votes < 0.6:
        return False
    for ln in lines:
        if fa_token_count(ln.text) >= 1:
            ln.text = reverse_line(ln.text)
    return True


# --- سرصفحه/پاورقی تکراری ---------------------------------------------------------------
PAGE_NUM = re.compile(r"^\s*[-–—|]?\s*(صفحه|ص\.?|page)?\s*\d+\s*(از|/|of)?\s*\d*\s*[-–—|]?\s*$", re.I)


def _line_key(t: str) -> str:
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", t.strip()))


def remove_repeated_lines(pages: list[PageResult]) -> int:
    n = len(pages)
    counts: Counter = Counter()
    for p in pages:
        counts.update({_line_key(ln.text) for ln in p.lines})
    threshold = max(3, int(0.5 * n))
    repeated = {k for k, c in counts.items() if n >= 3 and c >= threshold and len(k) < 100}
    removed = 0
    for p in pages:
        keep = [ln for ln in p.lines if not PAGE_NUM.match(ln.text) and _line_key(ln.text) not in repeated]
        removed += len(p.lines) - len(keep)
        p.lines = keep
    return removed


# --- جدول‌ها ----------------------------------------------------------------------------
def _cell(c) -> str:
    t = clean_fa(fix_line_direction(pre_clean(str(c or "")))[0]) if c else ""
    return t.replace("\n", " ").replace("|", "/")


def table_to_markdown(rows: list[list]) -> str | None:
    rows = [r for r in rows if r and any(c for c in r)]
    if len(rows) < 2 or max(len(r) for r in rows) < 2:
        return None
    width = max(len(r) for r in rows)
    rows = [[_cell(c) for c in r] + [""] * (width - len(r)) for r in rows]
    head, *body = rows
    lines = ["| " + " | ".join(head) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


# --- PDF ------------------------------------------------------------------------------
def extract_pdf(path: Path, opts: ExtractOptions, progress: Progress | None = None) -> ExtractResult:
    import pymupdf as fitz

    res = ExtractResult(pages=[])
    ocr_ok = ocr_status()["available"]
    if opts.method == "ocr" and not ocr_ok:
        raise ValueError("OCR در دسترس نیست: " + ocr_status().get("reason", ""))
    plumber_doc = None

    def plumber_page(i: int):
        nonlocal plumber_doc
        if plumber_doc is None:
            import pdfplumber
            plumber_doc = pdfplumber.open(str(path))
        return plumber_doc.pages[i]

    try:
        with fitz.open(str(path)) as doc:
            n = len(doc)
            no_text_pages: list[int] = []
            for i, fpage in enumerate(doc):
                pno = i + 1
                if progress:
                    progress(i / max(n, 1), f"استخراج صفحه {pno} از {n}")
                method = opts.method
                if method == "ocr":
                    lines = _ocr_lines(fpage, pno) if ocr_ok else []
                elif method == "pdfplumber":
                    lines, _ = _best_plumber(plumber_page(i), pno, opts.margin)
                else:
                    lines, method = _pymupdf_lines(fpage, pno, opts.margin), "pymupdf"
                    if opts.method == "auto":
                        raw = "\n".join(ln.text for ln in lines)
                        if len(raw.strip()) < 20 or looks_broken(raw):
                            if ocr_ok and opts.ocr_fallback:
                                lines, method = _ocr_lines(fpage, pno), "ocr"
                            else:
                                no_text_pages.append(pno)
                        elif _score_lines(lines) < 0.55:  # متن مشکوک: pdfplumber با بهترین x_tolerance
                            p_lines, p_score = _best_plumber(plumber_page(i), pno, opts.margin)
                            if p_score > _score_lines(lines) + 0.05:
                                lines, method = p_lines, "pdfplumber"

                raw_text = "\n".join(ln.text for ln in lines)
                page = PageResult(page=pno, lines=lines, method=method, stats=char_stats(raw_text))
                if opts.clean:
                    for ln in lines:
                        ln.text = pre_clean(ln.text)
                    if opts.fix_direction:
                        page.reversed = _fix_page_direction(lines)
                if opts.reverse_word_order:
                    for ln in lines:
                        ln.text = " ".join(reversed(ln.text.split()))
                page.vocab = vocab_score("\n".join(ln.text for ln in lines))
                res.pages.append(page)

                if opts.extract_tables:
                    try:
                        for t in plumber_page(i).extract_tables():
                            md = table_to_markdown(t)
                            if md:
                                res.tables.append({"page": pno, "markdown": md})
                    except Exception as e:  # pragma: no cover
                        log.warning("استخراج جدول صفحه %s ناموفق بود: %s", pno, e)

            if no_text_pages:
                hint = "" if ocr_ok else " (OCR در دسترس نیست؛ Tesseract با زبان fas را نصب کنید)"
                res.warnings.append(
                    f"{len(no_text_pages)} صفحه متن قابل استخراج نداشت یا فونت آن بدون ToUnicode بود{hint}: "
                    + ", ".join(map(str, no_text_pages[:20]))
                )
    finally:
        if plumber_doc is not None:
            plumber_doc.close()

    if opts.remove_repeated:
        removed = remove_repeated_lines(res.pages)
        if removed:
            res.warnings.append(f"{removed} خط سرصفحه/پاورقی/شماره صفحه حذف شد")
    if (rp := sum(p.reversed for p in res.pages)):
        res.warnings.append(f"جهت حروف در {rp} صفحه معکوس تشخیص داده و اصلاح شد")
    return res


# --- DOCX / TXT / MD ------------------------------------------------------------------
def extract_docx(path: Path, opts: ExtractOptions) -> ExtractResult:
    import docx

    d = docx.Document(str(path))
    lines: list[Line] = []
    for p in d.paragraphs:
        if not p.text.strip():
            lines.append(Line("", 1))  # خط خالی = مرز پاراگراف
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        heading = style.startswith(("heading", "title")) or "عنوان" in style
        text = pre_clean(p.text) if opts.clean else p.text
        lines.append(Line(text, 1, heading=heading))
        lines.append(Line("", 1))
    res = ExtractResult(pages=[PageResult(1, lines, "docx")])
    if opts.extract_tables:
        for t in d.tables:
            md = table_to_markdown([[c.text for c in row.cells] for row in t.rows])
            if md:
                res.tables.append({"page": 1, "markdown": md})
    res.pages[0].vocab = vocab_score("\n".join(ln.text for ln in lines))
    return res


def extract_text_file(path: Path, opts: ExtractOptions) -> ExtractResult:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1256"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    lines = []
    for t in text.splitlines():
        heading = path.suffix.lower() == ".md" and t.lstrip().startswith("#")
        if heading:
            t = t.lstrip("# ").strip()
        lines.append(Line(pre_clean(t) if opts.clean else t, 1, heading=heading))
    page = PageResult(1, lines, "text")
    page.vocab = vocab_score(text)
    return ExtractResult(pages=[page])


def extract_file(path: Path, opts: ExtractOptions | None = None, progress: Progress | None = None) -> ExtractResult:
    opts = opts or ExtractOptions()
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path, opts, progress)
    if ext == ".docx":
        return extract_docx(path, opts)
    if ext in (".txt", ".md"):
        return extract_text_file(path, opts)
    raise ValueError(f"نوع فایل پشتیبانی نمی‌شود: {ext}")


# --- مقایسه ابزارها برای آزمایشگاه (بخش ۳ راهنما) -------------------------------------
def compare_extractors(path: Path, page_no: int) -> dict:
    import pymupdf as fitz
    import pdfplumber

    out: dict = {"page": page_no, "methods": []}
    with fitz.open(str(path)) as doc:
        out["pages"] = len(doc)
        if not 1 <= page_no <= len(doc):
            raise ValueError("شماره صفحه نامعتبر است")
        fpage = doc[page_no - 1]
        candidates: list[tuple[str, str]] = [
            ("PyMuPDF", fpage.get_text("text", sort=True)),
        ]
        with pdfplumber.open(str(path)) as pdf:
            ppage = pdf.pages[page_no - 1]
            for tol in (1.5, 3.0):
                candidates.append((f"pdfplumber (x_tolerance={tol})", ppage.extract_text(x_tolerance=tol) or ""))
        if ocr_status()["available"]:
            candidates.append(("OCR (Tesseract)", "\n".join(ln.text for ln in _ocr_lines(fpage, page_no))))

    for name, raw in candidates:
        cleaned_lines = [pre_clean(t) for t in raw.splitlines()]
        tmp = [Line(t, page_no) for t in cleaned_lines]
        rev = _fix_page_direction(tmp)
        cleaned = clean_fa("\n".join(ln.text for ln in tmp))
        out["methods"].append({
            "name": name,
            "raw": raw,
            "cleaned": cleaned,
            "stats": char_stats(raw),
            "broken": looks_broken(raw),
            "vocab_raw": round(vocab_score(raw), 3),
            "vocab_clean": round(vocab_score(cleaned), 3),
            "reversed": rev,
        })
    return out


def lines_as_dicts(lines: list[Line]) -> list[dict]:
    return [asdict(ln) for ln in lines]
