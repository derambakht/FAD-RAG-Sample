"""نسخه جدید همان سند: تشخیص پیش از بارگذاری (نام) و پس از پردازش (محتوا)، و جایگزینی امن نسخه قبلی."""
import json

import openpyxl
import pytest

from app.db import get_db
from app.rag.ingest import version_key
from tests.test_spreadsheet import SAMPLE_XLSX, _wait

QTY_BEFORE = 141  # مقدار فروش اولین کالای مگامال در فایل نمونه
QTY_AFTER = 1141


@pytest.mark.parametrize("a,b,same", [
    ("sales-report-sample.xlsx", "sales-report-sample (1).xlsx", True),
    ("sales-report-sample.xlsx", "Sales-Report-Sample - Copy.xlsx", True),
    ("sales-report-sample.xlsx", "sales-report-sample_v2.xlsx", True),
    ("sales-report-sample.xlsx", "sales-report-sample.pdf", True),
    ("گزارش فروش.xlsx", "گزارش فروش - اصلاح‌شده.xlsx", True),
    ("گزارش فروش.xlsx", "گزارش فروش نسخه ۲.xlsx", True),
    ("report-1404-07.xlsx", "report-1404-08.xlsx", False),  # گزارش دوره دیگر، نسخه دیگر نیست
    ("plan.pdf", "semifinal.pdf", False),
])
def test_version_key(a, b, same):
    assert (version_key(a) == version_key(b)) is same


@pytest.fixture()
def kb(client):
    col = client.post("/api/collections", json={"name": f"نسخه‌ها {id(object())}", "embedding_provider": "openai",
                                                "embedding_model": "fake", "vector_backend": "chroma"}).json()
    yield col
    client.delete(f"/api/collections/{col['id']}")


def _edited_copy(tmp_path, name):
    """همان گزارش که کاربر یک ردیفش را اصلاح کرده است."""
    wb = openpyxl.load_workbook(SAMPLE_XLSX)
    ws = wb.active
    assert ws["D4"].value == QTY_BEFORE
    ws["D4"] = QTY_AFTER
    path = tmp_path / name
    wb.save(path)
    return path


def _upload(client, kb, path, name=None, replace=None):
    with open(path, "rb") as f:
        r = client.post(f"/api/collections/{kb['id']}/documents", files=[("files", (name or path.name, f))],
                        data={"replace": json.dumps(replace or {})})
    assert r.status_code == 202, r.text
    body = r.json()
    return body, (_wait(client, body["jobs"][0]["id"]) if body["jobs"] else None)


def _qty(kb):
    return sum(r["qty"] for r in get_db().sales_rows(kb["id"]))


def test_check_finds_previous_version_by_name(client, kb):
    _, job = _upload(client, kb, SAMPLE_XLSX)
    old = job["meta"]["document_id"]
    r = client.post(f"/api/collections/{kb['id']}/documents/check",
                    json={"files": ["sales-report-sample (1).xlsx", "other-report.xlsx"]}).json()["files"]
    assert [p["id"] for p in r[0]["previous"]] == [old] and r[1]["previous"] == []
    prev = r[0]["previous"][0]
    assert prev["rows"] == 279 and prev["branches"] == 3 and prev["status"] == "ready"


def test_replace_deletes_previous_only_after_new_version_is_ready(client, kb, tmp_path):
    _, job = _upload(client, kb, SAMPLE_XLSX)
    old = job["meta"]["document_id"]
    total = _qty(kb)
    edited = _edited_copy(tmp_path, "sales-report-sample.xlsx")
    _, job = _upload(client, kb, edited, replace={"sales-report-sample.xlsx": [old]})
    assert job["status"] == "done", job
    assert job["result"]["replaced"] == ["sales-report-sample.xlsx"] and job["result"]["overlaps"] == []
    docs = client.get(f"/api/collections/{kb['id']}/documents").json()
    assert [d["id"] for d in docs] == [job["meta"]["document_id"]]
    # فقط نسخه اصلاح‌شده حساب می‌شود: نه دوبار، نه مقدار قدیمی
    assert _qty(kb) == total - QTY_BEFORE + QTY_AFTER
    assert client.get(f"/api/collections/{kb['id']}").json()["table_rows"] == 279


def test_failed_new_version_keeps_previous(client, kb, tmp_path):
    _, job = _upload(client, kb, SAMPLE_XLSX)
    old = job["meta"]["document_id"]
    broken = tmp_path / "sales-report-sample.xlsx"
    broken.write_bytes(b"not an excel file")
    _, job = _upload(client, kb, broken, replace={"sales-report-sample.xlsx": [old]})
    assert job["status"] == "error"
    assert get_db().get_document(old)["status"] == "ready" and _qty(kb) > 0


def test_keep_both_flags_overlap_and_shows_double_count(client, kb, tmp_path):
    _, job = _upload(client, kb, SAMPLE_XLSX)
    old = job["meta"]["document_id"]
    total = _qty(kb)
    # نام دیگر، پس پیش از بارگذاری تشخیص داده نمی‌شود؛ پس از پردازش از روی سطرهای مشترک تشخیص داده می‌شود
    _, job = _upload(client, kb, _edited_copy(tmp_path, "فروش مهر.xlsx"))
    ov = job["result"]["overlaps"]
    assert [o["document_id"] for o in ov] == [old] and ov[0]["ratio"] == 1.0 and ov[0]["shared"] == ov[0]["total"]
    doc = get_db().get_document(job["meta"]["document_id"])
    assert doc["stats"]["overlaps"][0]["document_id"] == old
    assert _qty(kb) == 2 * total - QTY_BEFORE + QTY_AFTER  # همان خطری که در پنجره هشدار توضیح داده می‌شود


def test_processing_version_not_counted(client, kb):
    _, job = _upload(client, kb, SAMPLE_XLSX)
    did = job["meta"]["document_id"]
    db = get_db()
    db.update_document(did, status="processing")
    try:
        assert db.sales_rows(kb["id"]) == []
    finally:
        db.update_document(did, status="ready")


def test_text_document_overlap_detected(client, kb, tmp_path):
    text = "".join(f"# بخش {i}\n" + " ".join(f"جمله {j} از بند {i} آیین‌نامه درباره موضوع شماره {i} و جزئیات اجرایی آن است."
                                             for j in range(40)) + "\n\n" for i in range(8))
    a = tmp_path / "rules.md"
    a.write_text(text, encoding="utf-8")
    _, job = _upload(client, kb, a)
    old = job["meta"]["document_id"]
    b = tmp_path / "rules-2.md"
    b.write_text(text.replace("بند 3", "بند سه (اصلاح‌شده)"), encoding="utf-8")
    _, job = _upload(client, kb, b)
    assert [o["document_id"] for o in job["result"]["overlaps"]] == [old]


def test_replace_rejects_foreign_document(client, kb):
    other = client.post("/api/collections", json={"name": "دیگری برای نسخه", "embedding_provider": "openai",
                                                  "embedding_model": "fake", "vector_backend": "chroma"}).json()
    _, job = _upload(client, other, SAMPLE_XLSX)
    with open(SAMPLE_XLSX, "rb") as f:
        r = client.post(f"/api/collections/{kb['id']}/documents", files=[("files", (SAMPLE_XLSX.name, f))],
                        data={"replace": json.dumps({SAMPLE_XLSX.name: [job["meta"]["document_id"]]})})
    assert r.status_code == 404
    client.delete(f"/api/collections/{other['id']}")
